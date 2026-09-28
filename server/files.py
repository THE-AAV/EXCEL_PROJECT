"""The Files section: folders of files (Excel workbooks and anything else), like a simple file manager.

Every upload of a file is kept as a version (who, when, size, the sheets inside), so nothing is overwritten and
the audit log shows who uploaded, renamed, moved or deleted what. Deleting moves a file to "Recently deleted",
from where it can be brought back (an admin can delete it for good).

The file contents are kept by a Blobs store: in the data folder (the office computer), or in a private Hugging Face
dataset (the free online set-up, whose own disk is wiped on every restart). Big files are sent in pieces, so an
upload of up to MAX_FILE_MB works through any proxy and shows its progress.
"""
from __future__ import annotations

import hashlib
import logging
import os
import re
import shutil
import threading
import time
import uuid
from datetime import date, datetime
from pathlib import Path

from fastapi import HTTPException, status
from sqlalchemy import func, select, update

from . import db, store

log = logging.getLogger(__name__)
MAX_FILE_MB = int(os.environ.get("MAX_FILE_MB", "1024"))
CHUNK_MB = 8
PREVIEW_ROWS = 200
EXCEL = (".xlsx", ".xlsm", ".xltx", ".xltm")


# ------------------------------------------------------------------------------------------ where contents live
class LocalBlobs:
    """Files in <data folder>/files."""
    where = "on the computer running the app"

    def __init__(self, root: Path):
        self.root = root
        root.mkdir(parents=True, exist_ok=True)

    def put(self, key: str, src: Path) -> None:
        shutil.move(str(src), self.root / key)

    def path(self, key: str) -> Path:
        p = self.root / key
        if not p.exists():
            raise HTTPException(status.HTTP_404_NOT_FOUND, "This file's contents are missing")
        return p

    def delete(self, key: str) -> None:
        (self.root / key).unlink(missing_ok=True)


class HFBlobs:
    """Files in a private Hugging Face dataset (free, survives restarts), with a local copy as a cache."""
    where = "in your private Hugging Face storage"

    def __init__(self, repo: str, token: str, cache: Path):
        from huggingface_hub import HfApi
        self.repo, self.token, self.cache = repo, token, cache
        self.api = HfApi(token=token)
        cache.mkdir(parents=True, exist_ok=True)
        self.api.create_repo(repo, repo_type="dataset", private=True, exist_ok=True)

    def put(self, key: str, src: Path) -> None:
        self.api.upload_file(path_or_fileobj=str(src), path_in_repo=f"files/{key}", repo_id=self.repo,
                             repo_type="dataset", commit_message=f"Add {key}")
        shutil.move(str(src), self.cache / key)

    def path(self, key: str) -> Path:
        p = self.cache / key
        if not p.exists():
            from huggingface_hub import hf_hub_download
            got = hf_hub_download(self.repo, f"files/{key}", repo_type="dataset", token=self.token,
                                  local_dir=self.cache / ".hf")
            shutil.move(got, p)
        return p

    def delete(self, key: str) -> None:
        (self.cache / key).unlink(missing_ok=True)
        try:
            self.api.delete_file(f"files/{key}", repo_id=self.repo, repo_type="dataset", commit_message=f"Delete {key}")
        except Exception:   # already gone
            pass


class UnusableBlobs(LocalBlobs):
    """On a Hugging Face Space without working file storage: the Space's own disk is wiped on every restart,
    so uploads are refused (with the reason) rather than lost later."""

    def __init__(self, root: Path, problem: str):
        super().__init__(root)
        self.problem = problem
        self.where = f"nowhere yet: {problem}"

    def put(self, key: str, src: Path) -> None:
        src.unlink(missing_ok=True)
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, f"Files can't be kept yet: {self.problem}")


def make_blobs(data_dir: Path):
    repo, token = os.environ.get("HF_FILES_REPO", "").strip(), os.environ.get("HF_TOKEN", "").strip()
    if repo and token:
        try:
            return HFBlobs(repo, token, data_dir / "files-cache")
        except Exception as e:
            log.error("Hugging Face file storage not usable: %s", e)
            return UnusableBlobs(data_dir / "files", "the Space's HF_TOKEN or HF_FILES_REPO setting is not right "
                                 f"({str(e).splitlines()[0][:150]})")
    if os.environ.get("SPACE_ID"):
        return UnusableBlobs(data_dir / "files", "add HF_TOKEN and HF_FILES_REPO in the Space's settings "
                             "(see the README, step 2.5)")
    return LocalBlobs(data_dir / "files")


# ------------------------------------------------------------------------------------------ helpers
def clean_name(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "", (name or "").strip()).strip(" .")
    if not name:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Please type a name")
    return name[:200]


def can_change(user) -> None:
    if not user.get("can_files"):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "You can look at files but not change them. Ask an admin.")


def _folder(conn, folder_id):
    if folder_id is None:
        return None
    f = conn.execute(select(db.folders).where(db.folders.c.id == folder_id)).first()
    if f is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Folder not found")
    return f


def _file(conn, file_id: int):
    f = conn.execute(select(db.files).where(db.files.c.id == file_id)).first()
    if f is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "File not found")
    return f


def _live_file(conn, name: str, folder_id, exclude: int | None = None):
    q = select(db.files).where(db.files.c.deleted_at.is_(None), func.lower(db.files.c.name) == name.lower(),
                               db.files.c.folder_id.is_(None) if folder_id is None else db.files.c.folder_id == folder_id)
    if exclude:
        q = q.where(db.files.c.id != exclude)
    return conn.execute(q).first()


def _cell(v):
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v != v:
        return None
    return v if isinstance(v, (int, float, str, bool)) or v is None else str(v)


def scan_sheets(path: Path, name: str) -> list[dict]:
    """The sheets in an Excel file, with their size, read without loading the whole workbook."""
    if not name.lower().endswith(EXCEL):
        return []
    import openpyxl
    fh = path.open("rb")   # a file object: openpyxl refuses a path without an Excel extension
    try:
        wb = openpyxl.load_workbook(fh, read_only=True, data_only=True)
    except Exception:
        fh.close()
        return []
    out = []
    try:
        for ws in wb.worksheets:
            rows, cols = ws.max_row, ws.max_column
            # the size saved in the file can be missing or wrong (e.g. "A1"): count, unless the file is huge
            if (rows is None or cols is None or rows <= 1 or cols <= 1) and path.stat().st_size < 30 * 1024 * 1024:
                ws.reset_dimensions()
                rows = cols = 0
                for r in ws.iter_rows():
                    rows += 1
                    cols = max(cols, len(r))
            out.append({"name": ws.title, "rows": rows, "cols": cols})
    finally:
        wb.close()
        fh.close()
    return out


def col_letter(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def preview_sheet(path: Path, sheet: str, limit: int = PREVIEW_ROWS) -> dict:
    """The first rows of one sheet, with Excel's column letters and row numbers."""
    import openpyxl
    fh = path.open("rb")
    wb = openpyxl.load_workbook(fh, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Sheet not found")
        ws = wb[sheet]
        ws.reset_dimensions()   # read what is there, not the size the file claims
        rows, width = [], 0
        for n, r in enumerate(ws.iter_rows(values_only=True), 1):
            if n > limit:
                break
            vals = [_cell(v) for v in r]
            while vals and vals[-1] in (None, ""):
                vals.pop()
            width = max(width, len(vals))
            rows.append({"_n": n, **{f"c{i}": v for i, v in enumerate(vals)}})
    finally:
        wb.close()
        fh.close()
    width = min(width, 150)
    return {"cols": [[f"c{i}", col_letter(i)] for i in range(width)], "rows": rows,
            "more": len(rows) == limit}


# ------------------------------------------------------------------------------------------ the store
class FileStore:
    def __init__(self, engine, data_dir: Path, blobs=None):
        self.engine = engine
        self.blobs = blobs or make_blobs(data_dir)
        self.partial = data_dir / "partial-uploads"
        self.partial.mkdir(parents=True, exist_ok=True)
        self.uploads: dict[str, dict] = {}
        self.lock = threading.Lock()

    # ---- listing
    def tree(self) -> dict:
        with self.engine.connect() as conn:
            names = {r.id: r.username for r in conn.execute(select(db.users.c.id, db.users.c.username))}
            latest = select(db.file_versions.c.file_id, func.max(db.file_versions.c.number).label("n"),
                            func.count().label("count")).group_by(db.file_versions.c.file_id).subquery()
            vers = {r.file_id: r for r in conn.execute(
                select(db.file_versions, latest.c.count)
                .join(latest, (latest.c.file_id == db.file_versions.c.file_id) & (latest.c.n == db.file_versions.c.number)))}
            files = []
            for f in conn.execute(select(db.files).order_by(func.lower(db.files.c.name))):
                v = vers.get(f.id)
                files.append({"id": f.id, "name": f.name, "folder_id": f.folder_id, "created_at": f.created_at,
                              "created_by": names.get(f.created_by), "deleted_at": f.deleted_at,
                              "deleted_by": names.get(f.deleted_by),
                              "version": v.number if v else 0, "versions": v.count if v else 0,
                              "version_id": v.id if v else None, "size": v.size if v else 0,
                              "sheets": v.sheets if v else [], "changed_at": v.uploaded_at if v else f.created_at,
                              "changed_by": names.get(v.uploaded_by) if v else None})
            folders = [{"id": f.id, "name": f.name, "parent_id": f.parent_id, "created_at": f.created_at,
                        "created_by": names.get(f.created_by)}
                       for f in conn.execute(select(db.folders).order_by(func.lower(db.folders.c.name)))]
            master = conn.execute(select(db.import_batches).where(db.import_batches.c.status == "active")).first()
            count = conn.execute(select(func.count()).select_from(db.import_batches)
                                 .where(db.import_batches.c.status.in_(("active", "replaced", "undone")))).scalar_one()
        return {"folders": folders, "files": files, "where": self.blobs.where,
                "master": None if master is None else {
                    "changed_at": master.confirmed_at or master.created_at, "changed_by": names.get(master.created_by),
                    "last_change": master.filename, "versions": count}}

    def versions(self, file_id: int) -> list[dict]:
        with self.engine.connect() as conn:
            _file(conn, file_id)
            names = {r.id: r.username for r in conn.execute(select(db.users.c.id, db.users.c.username))}
            return [{"id": v.id, "number": v.number, "size": v.size, "sheets": v.sheets, "note": v.note,
                     "by": names.get(v.uploaded_by), "at": v.uploaded_at}
                    for v in conn.execute(select(db.file_versions).where(db.file_versions.c.file_id == file_id)
                                          .order_by(db.file_versions.c.number.desc()))]

    def activity(self, file_id: int) -> list[dict]:
        with self.engine.connect() as conn:
            rows = conn.execute(select(db.audit_log).where(db.audit_log.c.action.like("file.%"))
                                .order_by(db.audit_log.c.id.desc()).limit(5000))
            return [{"at": r.at, "by": r.username, "action": r.action, "details": r.details}
                    for r in rows if (r.details or {}).get("file_id") == file_id]

    def version(self, version_id: int):
        with self.engine.connect() as conn:
            v = conn.execute(select(db.file_versions, db.files.c.name)
                             .join(db.files, db.files.c.id == db.file_versions.c.file_id)
                             .where(db.file_versions.c.id == version_id)).first()
        if v is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "File not found")
        return v

    # ---- uploading, in pieces
    def start(self, user, name: str, size: int, folder_id=None, file_id=None) -> dict:
        can_change(user)
        name = clean_name(name)
        if size > MAX_FILE_MB * 1024 * 1024:
            raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, f"Files can be up to {MAX_FILE_MB:,} MB")
        free = shutil.disk_usage(self.partial).free
        if size * 2 > free:
            raise HTTPException(status.HTTP_507_INSUFFICIENT_STORAGE, "There is not enough free space for this file")
        with self.engine.connect() as conn:
            if file_id is not None:
                f = _file(conn, file_id)
                if f.deleted_at:
                    raise HTTPException(status.HTTP_409_CONFLICT, "Bring the file back first")
                folder_id = f.folder_id
            else:
                _folder(conn, folder_id)
        self._clean_stale()
        uid = uuid.uuid4().hex
        path = self.partial / uid
        path.write_bytes(b"")
        with self.lock:
            self.uploads[uid] = {"user": user["id"], "name": name, "size": size, "folder_id": folder_id,
                                 "file_id": file_id, "path": path, "started": time.time()}
        return {"upload_id": uid, "chunk_size": CHUNK_MB * 1024 * 1024}

    def _get_upload(self, user, uid: str) -> dict:
        up = self.uploads.get(uid)
        if up is None or up["user"] != user["id"]:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "This upload has expired. Please upload the file again.")
        return up

    async def chunk(self, user, uid: str, offset: int, stream) -> dict:
        up = self._get_upload(user, uid)
        have = up["path"].stat().st_size
        if offset != have:
            raise HTTPException(status.HTTP_409_CONFLICT, f"Expected the piece starting at {have}")
        written = 0
        with up["path"].open("ab") as fh:
            async for part in stream:
                written += len(part)
                if have + written > up["size"] or written > (CHUNK_MB + 1) * 1024 * 1024:
                    raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "More data than the file's size")
                fh.write(part)
        return {"received": have + written}

    def finish(self, user, uid: str, note: str = "") -> dict:
        up = self._get_upload(user, uid)
        path: Path = up["path"]
        if path.stat().st_size != up["size"]:
            raise HTTPException(status.HTTP_409_CONFLICT, "The upload is not complete")
        with self.lock:
            self.uploads.pop(uid, None)
        digest = hashlib.sha256()
        with path.open("rb") as fh:
            for block in iter(lambda: fh.read(1024 * 1024), b""):
                digest.update(block)
        sheets = scan_sheets(path, up["name"])
        key = f"{uuid.uuid4().hex}{Path(up['name']).suffix.lower()[:10]}"
        try:
            self.blobs.put(key, path)
        except HTTPException:
            raise
        except Exception as e:
            path.unlink(missing_ok=True)
            raise HTTPException(status.HTTP_502_BAD_GATEWAY, f"The file could not be stored ({e})")
        with self.engine.begin() as conn:
            file_id = up["file_id"]
            if file_id is None:   # the same name in the same folder: a new version of that file
                same = _live_file(conn, up["name"], up["folder_id"])
                file_id = same.id if same else None
            new = file_id is None
            if new:
                file_id = conn.execute(db.files.insert().values(
                    name=up["name"], folder_id=up["folder_id"], created_by=user["id"],
                    created_at=db.utcnow())).inserted_primary_key[0]
            number = (conn.execute(select(func.max(db.file_versions.c.number))
                                   .where(db.file_versions.c.file_id == file_id)).scalar() or 0) + 1
            vid = conn.execute(db.file_versions.insert().values(
                file_id=file_id, number=number, size=up["size"], sha256=digest.hexdigest(), stored_as=key,
                sheets=sheets, note=(note or ("Uploaded" if new else "New version"))[:255], uploaded_by=user["id"],
                uploaded_at=db.utcnow())).inserted_primary_key[0]
            name = _file(conn, file_id).name
            store.audit(conn, user, "file.upload" if new else "file.version", file_id=file_id, file=name,
                        version=number, size=up["size"], sheets=[s["name"] for s in sheets])
        return {"file_id": file_id, "version_id": vid, "version": number, "name": name, "new": new}

    def cancel(self, user, uid: str) -> None:
        up = self._get_upload(user, uid)
        with self.lock:
            self.uploads.pop(uid, None)
        up["path"].unlink(missing_ok=True)

    def _clean_stale(self) -> None:
        cutoff = time.time() - 24 * 3600
        with self.lock:
            for uid in [u for u, v in self.uploads.items() if v["started"] < cutoff]:
                self.uploads.pop(uid)["path"].unlink(missing_ok=True)
        live = {v["path"].name for v in self.uploads.values()}
        for p in self.partial.iterdir():
            if p.name not in live and p.stat().st_mtime < cutoff:
                p.unlink(missing_ok=True)

    # ---- folders
    def add_folder(self, user, name: str, parent_id=None) -> dict:
        can_change(user)
        name = clean_name(name)
        with self.engine.begin() as conn:
            _folder(conn, parent_id)
            if self._folder_named(conn, name, parent_id):
                raise HTTPException(status.HTTP_409_CONFLICT, f"There is already a folder called {name} here")
            fid = conn.execute(db.folders.insert().values(name=name, parent_id=parent_id, created_by=user["id"],
                                                          created_at=db.utcnow())).inserted_primary_key[0]
            store.audit(conn, user, "file.folder_add", folder_id=fid, folder=name)
        return {"id": fid}

    @staticmethod
    def _folder_named(conn, name, parent_id, exclude=None):
        q = select(db.folders.c.id).where(func.lower(db.folders.c.name) == name.lower(),
                                          db.folders.c.parent_id.is_(None) if parent_id is None
                                          else db.folders.c.parent_id == parent_id)
        if exclude:
            q = q.where(db.folders.c.id != exclude)
        return conn.execute(q).first()

    def edit_folder(self, user, folder_id: int, name=None, parent_id=..., ) -> None:
        can_change(user)
        with self.engine.begin() as conn:
            f = _folder(conn, folder_id)
            values, changes = {}, {}
            new_parent = f.parent_id if parent_id is ... else parent_id
            if parent_id is not ... and parent_id != f.parent_id:
                p = parent_id
                while p is not None:   # not into itself or one of its own subfolders
                    if p == folder_id:
                        raise HTTPException(status.HTTP_409_CONFLICT, "A folder can't be moved into itself")
                    p = _folder(conn, p).parent_id
                values["parent_id"] = parent_id
                changes["moved"] = True
            if name is not None and clean_name(name) != f.name:
                values["name"] = changes["renamed_to"] = clean_name(name)
            if values:
                if self._folder_named(conn, values.get("name", f.name), new_parent, exclude=folder_id):
                    raise HTTPException(status.HTTP_409_CONFLICT, "There is already a folder with that name there")
                conn.execute(update(db.folders).where(db.folders.c.id == folder_id).values(**values))
                store.audit(conn, user, "file.folder_edit", folder_id=folder_id, folder=f.name, **changes)

    def delete_folder(self, user, folder_id: int) -> None:
        can_change(user)
        with self.engine.begin() as conn:
            f = _folder(conn, folder_id)
            if conn.execute(select(db.folders.c.id).where(db.folders.c.parent_id == folder_id)).first() or \
                    conn.execute(select(db.files.c.id).where(db.files.c.folder_id == folder_id,
                                                              db.files.c.deleted_at.is_(None))).first():
                raise HTTPException(status.HTTP_409_CONFLICT, "Only an empty folder can be deleted: move or delete "
                                                              "what is in it first")
            # files deleted earlier from this folder go back to the top level if they are ever brought back
            conn.execute(update(db.files).where(db.files.c.folder_id == folder_id).values(folder_id=None))
            conn.execute(db.folders.delete().where(db.folders.c.id == folder_id))
            store.audit(conn, user, "file.folder_delete", folder_id=folder_id, folder=f.name)

    # ---- files
    def edit_file(self, user, file_id: int, name=None, folder_id=...) -> None:
        can_change(user)
        with self.engine.begin() as conn:
            f = _file(conn, file_id)
            values, changes = {}, {}
            if folder_id is not ... and folder_id != f.folder_id:
                _folder(conn, folder_id)
                values["folder_id"] = folder_id
                changes["moved_to"] = _folder(conn, folder_id).name if folder_id else "Files (top level)"
            if name is not None and clean_name(name) != f.name:
                values["name"] = changes["renamed_to"] = clean_name(name)
            if values:
                if f.deleted_at is None and _live_file(conn, values.get("name", f.name),
                                                       values.get("folder_id", f.folder_id), exclude=file_id):
                    raise HTTPException(status.HTTP_409_CONFLICT, "There is already a file with that name there")
                conn.execute(update(db.files).where(db.files.c.id == file_id).values(**values))
                store.audit(conn, user, "file.edit", file_id=file_id, file=f.name, **changes)

    def delete_file(self, user, file_id: int) -> None:
        can_change(user)
        with self.engine.begin() as conn:
            f = _file(conn, file_id)
            if f.deleted_at is None:
                conn.execute(update(db.files).where(db.files.c.id == file_id)
                             .values(deleted_at=db.utcnow(), deleted_by=user["id"]))
                store.audit(conn, user, "file.delete", file_id=file_id, file=f.name)

    def restore_file(self, user, file_id: int) -> None:
        can_change(user)
        with self.engine.begin() as conn:
            f = _file(conn, file_id)
            if f.deleted_at is None:
                return
            if _live_file(conn, f.name, f.folder_id, exclude=file_id):
                raise HTTPException(status.HTTP_409_CONFLICT, f"There is already a file called {f.name} there: "
                                                              "rename or move that one first")
            conn.execute(update(db.files).where(db.files.c.id == file_id).values(deleted_at=None, deleted_by=None))
            store.audit(conn, user, "file.restore", file_id=file_id, file=f.name)

    def purge_file(self, user, file_id: int) -> None:
        """Deletes a file from Recently deleted for good (admins), freeing its space."""
        with self.engine.begin() as conn:
            f = _file(conn, file_id)
            if f.deleted_at is None:
                raise HTTPException(status.HTTP_409_CONFLICT, "Delete the file first")
            keys = [v.stored_as for v in conn.execute(select(db.file_versions.c.stored_as)
                                                       .where(db.file_versions.c.file_id == file_id))]
            conn.execute(db.file_versions.delete().where(db.file_versions.c.file_id == file_id))
            conn.execute(db.files.delete().where(db.files.c.id == file_id))
            store.audit(conn, user, "file.purge", file_id=file_id, file=f.name, versions=len(keys))
        for k in keys:
            self.blobs.delete(k)
