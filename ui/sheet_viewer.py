"""Excel-like read-only view of any workbook: sheet tabs, column letters, row numbers, the workbook's
own fonts, fills, borders, number formats, merged cells and frozen panes."""
from __future__ import annotations

import colorsys
import html
import io
import re
import warnings
from datetime import date, datetime, time

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter
from lxml import etree

DEFAULT_THEME = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5",
                 "70AD47", "0563C1", "954F72"]
BORDER_CSS = {"thin": "1px solid", "hair": "1px dotted", "dotted": "1px dotted", "dashed": "1px dashed",
              "medium": "2px solid", "mediumDashed": "2px dashed", "thick": "3px solid", "double": "3px double",
              "dashDot": "1px dashed", "mediumDashDot": "2px dashed", "slantDashDot": "2px dashed",
              "dashDotDot": "1px dashed", "mediumDashDotDot": "2px dashed"}
MAX_CELLS = 400_000


# --------------------------------------------------------------------------- colours
def _theme_colors(wb) -> list[str]:
    try:
        root = etree.fromstring(wb.loaded_theme)
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        scheme = root.find(".//a:clrScheme", ns)
        vals = []
        for child in scheme:
            c = child[0]
            vals.append(c.get("lastClr") or c.get("val"))
        # Excel's theme index order swaps the first two pairs
        return [vals[1], vals[0], vals[3], vals[2], *vals[4:]]
    except Exception:
        return DEFAULT_THEME


def _tint(hex6: str, tint: float) -> str:
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return "".join(f"{round(x * 255):02X}" for x in (r, g, b))


def _color(c, theme) -> str | None:
    if c is None:
        return None
    try:
        if c.type == "rgb" and isinstance(c.rgb, str):
            v = c.rgb[-6:]
            if c.rgb.startswith("00") and v == "000000" and len(c.rgb) == 8:
                return None  # 'auto' / unset
        elif c.type == "theme":
            v = theme[c.theme] if c.theme < len(theme) else None
        elif c.type == "indexed":
            if c.indexed in (64, 65):
                return None
            v = COLOR_INDEX[c.indexed][-6:]
        else:
            return None
        if v is None:
            return None
        if c.tint:
            v = _tint(v, c.tint)
        return "#" + v
    except Exception:
        return None


# --------------------------------------------------------------------------- number formats
def _sections(code: str) -> list[str]:
    out, cur, q = [], "", False
    for ch in code:
        if ch == '"':
            q = not q
        if ch == ";" and not q:
            out.append(cur)
            cur = ""
        else:
            cur += ch
    return out + [cur]


def group(num: str, indian: bool) -> str:
    if len(num) <= 3:
        return num
    head, tail = num[:-3], num[-3:]
    if indian:
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        return ",".join(parts) + "," + tail
    return f"{int(num):,}"


def _general(v: float) -> str:
    if float(v).is_integer() and abs(v) < 1e15:
        return str(int(v))
    s = f"{v:.10g}"
    return s


_DATE_TOKENS = re.compile(r"yyyy|yy|mmmmm|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|AM/PM|am/pm|A/P", re.I)


def _is_date_code(code: str) -> bool:
    clean = re.sub(r'"[^"]*"|\\.|\[[^\]]*\]', "", code)
    return bool(re.search(r"[dy]", clean, re.I)) or bool(re.search(r"[mhs]", clean, re.I)
                                                          and not re.search(r"[0#?]", clean))


def _format_date(v, code: str) -> str:
    code = _sections(code)[0]
    code = re.sub(r"\[[^\]]*\]", "", code)
    out, i, last = [], 0, None
    tokens = list(_DATE_TOKENS.finditer(code))
    for n, m in enumerate(tokens):
        lit = code[i:m.start()]
        out.append(re.sub(r'"([^"]*)"', r"\1", lit).replace("\\", ""))
        t = m.group(0)
        tl = t.lower()
        nxt = tokens[n + 1].group(0).lower() if n + 1 < len(tokens) else ""
        minute = tl in ("m", "mm") and (last in ("h", "hh") or nxt in ("s", "ss"))
        if tl == "yyyy":
            out.append(f"{v.year:04d}")
        elif tl == "yy":
            out.append(f"{v.year % 100:02d}")
        elif minute:
            out.append(f"{getattr(v, 'minute', 0):0{len(tl)}d}")
        elif tl == "mmmmm":
            out.append(v.strftime("%B")[0])
        elif tl == "mmmm":
            out.append(v.strftime("%B"))
        elif tl == "mmm":
            out.append(v.strftime("%b"))
        elif tl == "mm":
            out.append(f"{v.month:02d}")
        elif tl == "m":
            out.append(str(v.month))
        elif tl == "dddd":
            out.append(v.strftime("%A"))
        elif tl == "ddd":
            out.append(v.strftime("%a"))
        elif tl == "dd":
            out.append(f"{v.day:02d}")
        elif tl == "d":
            out.append(str(v.day))
        elif tl in ("hh", "h"):
            h = getattr(v, "hour", 0)
            if re.search(r"am/pm|a/p", code, re.I):
                h = h % 12 or 12
            out.append(f"{h:02d}" if tl == "hh" else str(h))
        elif tl in ("ss", "s"):
            out.append(f"{getattr(v, 'second', 0):0{len(tl)}d}")
        elif tl in ("am/pm", "a/p"):
            out.append("AM" if getattr(v, "hour", 0) < 12 else "PM")
        last = tl
        i = m.end()
    out.append(re.sub(r'"([^"]*)"', r"\1", code[i:]).replace("\\", ""))
    return "".join(out).strip()


def format_value(v, code: str, indian: bool = True) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (datetime, date)):
        if code in ("General", "", None) or not _is_date_code(code):
            code = "dd-mm-yyyy" if not isinstance(v, datetime) or (v.hour, v.minute) == (0, 0) else "dd-mm-yyyy hh:mm"
        return _format_date(v, code)
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if not isinstance(v, (int, float)):
        return str(v)
    if not code or code == "General":
        return _general(v)
    if code == "@":
        return _general(v)
    secs = _sections(code)
    neg_sec = v < 0 and len(secs) > 1
    sec = secs[1] if neg_sec else secs[2] if v == 0 and len(secs) > 2 else secs[0]
    x = abs(v) if neg_sec else v
    sec = re.sub(r"\[[^\]]*\]", "", sec)          # colours / conditions
    sec = re.sub(r"_.", "", sec)                   # padding
    sec = re.sub(r"\*.", "", sec)                  # fill characters
    placeholders = re.search(r"[0#?]", re.sub(r'"[^"]*"', "", sec))
    if not placeholders:
        return re.sub(r'"([^"]*)"', r"\1", sec).replace("\\", "").strip() or (_general(v) if sec.strip() == "" else "")
    stripped = re.sub(r'"[^"]*"', lambda m: "\x00" * len(m.group(0)), sec)
    first = re.search(r"[0#?]", stripped).start()
    last_m = [m.end() for m in re.finditer(r"[0#?]", stripped)][-1]
    prefix = re.sub(r'"([^"]*)"', r"\1", sec[:first]).replace("\\", "")
    suffix = re.sub(r'"([^"]*)"', r"\1", sec[last_m:]).replace("\\", "")
    body = sec[first:last_m]
    if "%" in sec:
        x = x * 100 ** sec.count("%")
    decimals = len(re.findall(r"[0#?]", body.split(".", 1)[1])) if "." in body else 0
    s = f"{abs(x):.{decimals}f}"
    whole, _, frac = s.partition(".")
    if "," in body:
        whole = group(whole, indian)
    if not re.search(r"0", body.split(".")[0]) and whole == "0" and frac:
        whole = ""
    txt = whole + ("." + frac if frac else "")
    if x == 0 and not re.search(r"0", body):   # '?' / '#' placeholders show nothing for zero
        txt = ""
    if x < 0 and not neg_sec:
        txt = "-" + txt
    return (prefix + txt + suffix).strip()


# --------------------------------------------------------------------------- building a sheet model
def load_workbook_views(source, indian: bool = True) -> list[dict]:
    """Read a workbook (path or bytes) into plain dicts that are cheap to render."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        data = io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source
        wb_v = openpyxl.load_workbook(data, data_only=True)
        if isinstance(data, io.BytesIO):
            data.seek(0)
        wb_f = openpyxl.load_workbook(data, data_only=False)
    theme = _theme_colors(wb_f)
    try:
        from mis_reports.xlsx_edit import Book
        book = Book(source if isinstance(source, (bytes, bytearray)) else str(source))
        uncalc = {name: book.sheet(name).uncalculated() for name in book.sheet_parts}
    except Exception:
        uncalc = None
    active = wb_f.active.title if wb_f.active is not None else None
    views = []
    for ws in wb_f.worksheets:
        if ws.sheet_state == "veryHidden":
            continue
        wv = wb_v[ws.title]
        pending = uncalc.get(ws.title, set()) if uncalc is not None else None
        views.append({**_sheet_view(ws, wv, theme, indian, pending), "active": ws.title == active})
    return views


def _sheet_view(ws, wv, theme, indian, uncalculated=None) -> dict:
    max_r, max_c = 0, 0
    for row in wv.iter_rows():
        for c in row:
            if c.value is not None:
                max_r, max_c = max(max_r, c.row), max(max_c, c.column)
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                max_r, max_c = max(max_r, c.row), max(max_c, c.column)
    for m in ws.merged_cells.ranges:
        max_r, max_c = max(max_r, m.max_row), max(max_c, m.max_col)
    if max_r * max_c > MAX_CELLS:
        max_r = max(1, MAX_CELLS // max(max_c, 1))

    default_w = ws.sheet_format.defaultColWidth or ws.sheet_format.baseColWidth or 8.43
    widths, hidden_cols = {}, set()
    for key, dim in ws.column_dimensions.items():
        lo, hi = dim.min or 0, dim.max or 0
        if not lo:
            continue
        for ci in range(lo, min(hi, max_c) + 1):
            if dim.width:
                widths[ci] = dim.width
            if dim.hidden:
                hidden_cols.add(ci)
    default_h = ws.sheet_format.defaultRowHeight or 15
    heights, hidden_rows = {}, set()
    for ri, dim in ws.row_dimensions.items():
        if ri > max_r:
            continue
        if dim.ht:
            heights[ri] = dim.ht
        if dim.hidden:
            hidden_rows.add(ri)

    merged = {}
    covered = set()
    for m in ws.merged_cells.ranges:
        merged[(m.min_row, m.min_col)] = (m.max_row - m.min_row + 1, m.max_col - m.min_col + 1)
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if (r, c) != (m.min_row, m.min_col):
                    covered.add((r, c))

    styles, style_ids, cells = [], {}, {}
    for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c):
        for c in row:
            if (c.row, c.column) in covered:
                continue
            value = wv.cell(c.row, c.column).value
            formula = c.value if isinstance(c.value, str) and c.value.startswith("=") else None
            if formula is None and hasattr(c.value, "text"):   # array formulas
                formula = "=" + str(c.value.text).lstrip("=")
            css = _cell_css(c, value, theme)
            if not css and value is None and formula is None:
                continue
            sid = style_ids.setdefault(css, len(style_ids))
            if sid == len(styles):
                styles.append(css)
            text = format_value(value, c.number_format, indian)
            if value is None and formula and (uncalculated is None or c.coordinate in uncalculated):
                text, pending = "", True
            else:
                pending = False
            cells[(c.row, c.column)] = {"t": text, "s": sid, "f": formula, "p": pending,
                                        "raw": "" if value is None else str(value)}
    af = None
    if ws.auto_filter is not None and ws.auto_filter.ref:
        try:
            from openpyxl.utils.cell import range_boundaries
            c1, r1, c2, _ = range_boundaries(ws.auto_filter.ref)
            af = (r1, c1, c2)
        except Exception:
            af = None
    fr, fc = 0, 0
    if ws.freeze_panes:
        m = re.match(r"([A-Z]+)(\d+)", ws.freeze_panes)
        fc = openpyxl.utils.column_index_from_string(m.group(1)) - 1
        fr = int(m.group(2)) - 1
    return {"name": ws.title, "state": ws.sheet_state, "rows": max_r, "cols": max_c, "widths": widths,
            "default_w": default_w, "heights": heights, "default_h": default_h, "hidden_rows": hidden_rows,
            "hidden_cols": hidden_cols, "merged": merged, "cells": cells, "styles": styles,
            "freeze": (fr, fc), "autofilter": af, "grid": ws.sheet_view.showGridLines is not False,
            "tab_color": _color(ws.sheet_properties.tabColor, theme) if ws.sheet_properties.tabColor else None}


def _cell_css(c, value, theme) -> str:
    parts = []
    f = c.font
    if f is not None:
        if f.b:
            parts.append("font-weight:700")
        if f.i:
            parts.append("font-style:italic")
        if f.u:
            parts.append("text-decoration:underline")
        if f.strike:
            parts.append("text-decoration:line-through")
        col = _color(f.color, theme)
        if col and col.upper() != "#000000":
            parts.append(f"color:{col}")
        if f.sz and abs(float(f.sz) - 11) > 0.1:
            parts.append(f"font-size:{float(f.sz):g}pt")
    fill = c.fill
    if fill is not None and fill.fill_type == "solid":
        col = _color(fill.fgColor, theme) or _color(fill.bgColor, theme)
        if col:
            parts.append(f"background:{col}")
    a = c.alignment
    h = a.horizontal if a is not None else None
    if h in ("center", "centerContinuous"):
        parts.append("text-align:center")
    elif h == "right":
        parts.append("text-align:right")
    elif h == "left":
        parts.append("text-align:left")
    elif isinstance(value, (int, float, datetime, date)) and not isinstance(value, bool):
        parts.append("text-align:right")
    elif isinstance(value, bool):
        parts.append("text-align:center")
    if a is not None and a.wrap_text:
        parts.append("white-space:normal")
    if a is not None and a.vertical in ("top", "center"):
        parts.append(f"vertical-align:{'top' if a.vertical == 'top' else 'middle'}")
    b = c.border
    if b is not None:
        for side in ("left", "right", "top", "bottom"):
            s = getattr(b, side)
            if s is not None and s.style:
                css = BORDER_CSS.get(s.style, "1px solid")
                parts.append(f"border-{side}:{css} {_color(s.color, theme) or '#000'}")
    return ";".join(parts)


# --------------------------------------------------------------------------- filter & sort (like Excel's AutoFilter)
BLANK = "(Blanks)"


def default_header_row(v: dict) -> int:
    if v.get("autofilter"):
        return v["autofilter"][0]
    return v["freeze"][0] or 1


def header_labels(v: dict, header_row: int) -> dict[int, str]:
    """Column number -> 'H · PARTY NAME' for columns that have data below the header."""
    used = {c for (r, c) in v["cells"] if r > header_row and v["cells"][(r, c)]["t"]}
    out = {}
    for c in sorted(used):
        if c in v["hidden_cols"]:
            continue
        head = v["cells"].get((header_row, c), {}).get("t", "")
        out[c] = f"{get_column_letter(c)} · {head}" if head else get_column_letter(c)
    return out


def data_rows(v: dict, header_row: int) -> list[int]:
    return [r for r in range(header_row + 1, v["rows"] + 1) if r not in v["hidden_rows"]]


def _num(cell) -> float | None:
    if cell is None:
        return None
    raw = cell["raw"]
    try:
        return float(raw)
    except ValueError:
        try:
            return datetime.fromisoformat(raw).timestamp() / 86400 + 25569   # dates as Excel day numbers
        except ValueError:
            return None


def date_number(d) -> float:
    """Same scale _num uses for dates, for min / max filters on date columns."""
    return datetime(d.year, d.month, d.day).timestamp() / 86400 + 25569


def is_date_column(v: dict, header_row: int, col: int) -> bool:
    cells = [v["cells"].get((r, col)) for r in data_rows(v, header_row)]
    filled = [c for c in cells if c and c["raw"]]
    dated = 0
    for c in filled:
        try:
            datetime.fromisoformat(c["raw"])
            dated += 1
        except ValueError:
            pass
    return bool(filled) and dated >= 0.8 * len(filled)


def column_values(v: dict, header_row: int, col: int) -> list[str]:
    """Distinct values shown in a column below the header, in Excel's checklist order."""
    vals, nums = set(), {}
    blank = False
    for r in data_rows(v, header_row):
        cell = v["cells"].get((r, col))
        t = cell["t"] if cell else ""
        if not t:
            blank = True
            continue
        vals.add(t)
        n = _num(cell)
        if n is not None:
            nums[t] = n
    ordered = sorted(vals, key=lambda t: (0, nums[t], "") if t in nums else (1, 0, t.lower()))
    return ordered + ([BLANK] if blank else [])


def is_numeric_column(v: dict, header_row: int, col: int) -> bool:
    cells = [v["cells"].get((r, col)) for r in data_rows(v, header_row)]
    filled = [c for c in cells if c and c["t"]]
    return bool(filled) and sum(_num(c) is not None for c in filled) >= 0.8 * len(filled)


def filter_rows(v: dict, header_row: int, filters: dict, sort: tuple | None = None) -> list[int]:
    """Rows below the header that pass every column filter, optionally sorted.

    filters: {col: {"values": [...], "text": str, "min": float, "max": float}}; empty parts are ignored.
    sort: (col, ascending)."""
    rows = []
    for r in data_rows(v, header_row):
        ok = True
        for col, f in filters.items():
            cell = v["cells"].get((r, col))
            t = cell["t"] if cell else ""
            if f.get("values") and (t or BLANK) not in f["values"]:
                ok = False
            elif f.get("text") and f["text"].lower() not in t.lower():
                ok = False
            elif f.get("min") is not None or f.get("max") is not None:
                n = _num(cell)
                if n is None or (f.get("min") is not None and n < f["min"]) or \
                        (f.get("max") is not None and n > f["max"]):
                    ok = False
            if not ok:
                break
        if ok:
            rows.append(r)
    if sort:
        col, asc = sort

        def key(r):
            cell = v["cells"].get((r, col))
            n = _num(cell)
            t = cell["t"] if cell else ""
            if not t:
                return (2, 0, "")                     # blanks always last, as in Excel
            return (0, n, "") if n is not None else (1, 0, t.lower())
        filled = [r for r in rows if key(r)[0] < 2]
        blanks = [r for r in rows if key(r)[0] == 2]
        rows = sorted(filled, key=key, reverse=not asc) + blanks
    return rows


# --------------------------------------------------------------------------- rendering
def col_px(width_chars: float) -> int:
    return int(round(width_chars * 7 + 5))


def render_html(v: dict, find: str = "", show_formulas: bool = False, zoom: int = 100,
                row_from: int = 1, row_to: int | None = None, height: int = 640,
                header_row: int | None = None, shown_rows: list[int] | None = None,
                filter_cols: set | None = None) -> tuple[str, int]:
    """Return (html, number of matching cells for find).

    header_row / shown_rows: when filtering, the rows up to header_row are always shown and below it only
    shown_rows, in that order. filter_cols: columns that have an active filter (marked on the header)."""
    fr, fc = v["freeze"]
    row_to = min(row_to or v["rows"], v["rows"])
    cols = [c for c in range(1, v["cols"] + 1) if c not in v["hidden_cols"]]
    widths = {c: col_px(v["widths"].get(c, v["default_w"])) for c in cols}
    if shown_rows is not None and header_row:
        rows = [r for r in range(1, header_row + 1) if r not in v["hidden_rows"]] + list(shown_rows)
    else:
        rows = [r for r in range(1, v["rows"] + 1) if r not in v["hidden_rows"] and (r <= fr or row_from <= r <= row_to)]
    keep_top = max(fr, header_row or 0)
    needle = find.strip().lower()
    hits = 0
    if needle:
        keep = set()
        for (r, c), cell in v["cells"].items():
            text = (cell["f"] or "") if show_formulas else cell["t"]
            if needle in text.lower() or needle in cell["raw"].lower():
                keep.add(r)
                hits += 1
        rows = [r for r in rows if r <= keep_top or r in keep]
    heights = {r: max(16, int(round(v["heights"].get(r, v["default_h"]) * 4 / 3))) for r in rows}
    head_h, rn_w = 22, 44

    # sticky offsets for frozen rows / columns
    top, acc = {}, head_h
    for r in rows:
        if r <= fr:
            top[r] = acc
            acc += heights[r]
    left, acc = {}, rn_w
    for c in cols:
        if c <= fc:
            left[c] = acc
            acc += widths[c]

    grid = "#e3e3e3" if v["grid"] else "transparent"
    css = [f".s{i}{{{s}}}" for i, s in enumerate(v["styles"])]
    out = [f"""<!doctype html><html><head><meta charset="utf-8"><style>
body{{margin:0;font-family:Calibri,Carlito,'Segoe UI',Arial,sans-serif;font-size:11pt;color:#000;background:#fff}}
.wrap{{height:{height}px;overflow:auto;border:1px solid #c8c8c8}}
table{{border-collapse:separate;border-spacing:0;table-layout:fixed;zoom:{zoom / 100};background:#fff;
       width:{rn_w + sum(widths.values())}px}}
td,th{{border-right:1px solid {grid};border-bottom:1px solid {grid};padding:1px 4px;white-space:nowrap;overflow:visible;
      box-sizing:border-box;vertical-align:bottom}}
td.clip,td.fz{{overflow:hidden;text-overflow:clip}}
th{{background:#f3f3f3;color:#444;font-weight:400;font-size:9.5pt;text-align:center;border-color:#cfcfcf;
    position:sticky;top:0;z-index:3}}
th.rn{{left:0;z-index:4;text-align:center}}
td.rn{{position:sticky;left:0;z-index:2;background:#f3f3f3;color:#444;font-size:9.5pt;text-align:center;
       border-color:#cfcfcf}}
.fz{{position:sticky;z-index:1;background:#fff}}
td.pending::after{{content:'…';color:#9a9a9a}}
td.hit{{outline:2px solid #f2b705;outline-offset:-2px}}
.fb{{float:right;font-size:8px;line-height:12px;border:1px solid #a6a6a6;background:#fff;color:#444;padding:0 2px;
     margin:1px 0 0 3px;border-radius:2px;font-weight:400;font-style:normal}}
.fb.on{{background:#2a78d6;border-color:#2a78d6;color:#fff}}
tr.sorted td.rn{{color:#2a78d6}}
{''.join(css)}
</style></head><body><div class="wrap"><table>"""]
    out.append("<colgroup>" + f'<col style="width:{rn_w}px">' +
               "".join(f'<col style="width:{widths[c]}px">' for c in cols) + "</colgroup>")
    out.append("<tr>" + '<th class="rn"></th>' + "".join(
        f'<th style="{"left:" + str(left[c]) + "px;z-index:4;" if c in left else ""}">{get_column_letter(c)}</th>'
        for c in cols) + "</tr>")
    marked = set()
    if header_row:
        af = v.get("autofilter")
        if af and af[0] == header_row:
            marked = set(range(af[1], af[2] + 1))
        else:
            marked = {c for (r, c), cell in v["cells"].items() if r == header_row and cell["t"]}
        marked |= set(filter_cols or ())
    skip = set()
    for r in rows:
        rstyle = f"height:{heights[r]}px"
        cells_html = [f'<td class="rn" style="{"top:" + str(top[r]) + "px;z-index:3;" if r in top else ""}">{r}</td>']
        for c in cols:
            if (r, c) in skip:
                continue
            cell = v["cells"].get((r, c))
            span = v["merged"].get((r, c))
            attrs, classes, style = [], [], []
            if span:
                rs, cs = span
                if cs > 1:
                    attrs.append(f'colspan="{sum(1 for x in range(c, c + cs) if x in widths)}"')
                if rs > 1:
                    attrs.append(f'rowspan="{sum(1 for x in range(r, r + rs) if x in heights)}"')
                for rr in range(r, r + rs):
                    for cc in range(c, c + cs):
                        if (rr, cc) != (r, c):
                            skip.add((rr, cc))
            if cell is not None:
                classes.append(f"s{cell['s']}")
                nxt = v["cells"].get((r, c + 1))
                if span or "white-space:normal" in v["styles"][cell["s"]] or (nxt and (nxt["t"] or nxt["p"])):
                    classes.append("clip")
                text = (cell["f"] or cell["t"]) if show_formulas else cell["t"]
                if cell["p"] and not show_formulas:
                    classes.append("pending")
                title = cell["f"] or (cell["t"] if len(cell["t"]) > 12 else "")
                if title:
                    attrs.append(f'title="{html.escape(title, quote=True)}"')
                if needle and (needle in text.lower() or needle in cell["raw"].lower()):
                    classes.append("hit")
            else:
                text = ""
            if r in top or c in left:
                classes.append("fz")
                if r in top:
                    style.append(f"top:{top[r]}px")
                if c in left:
                    style.append(f"left:{left[c]}px")
                if r in top and c in left:
                    style.append("z-index:2")
            if show_formulas and cell and cell["f"]:
                style.append("text-align:left;color:#1f4e79")
            cls = f' class="{" ".join(classes)}"' if classes else ""
            sty = f' style="{";".join(style)}"' if style else ""
            icon = ""
            if header_row and r == header_row and c in marked:
                icon = f'<span class="fb on">⏷</span>' if c in (filter_cols or set()) else '<span class="fb">▾</span>'
            cells_html.append(f"<td{cls}{sty}{' ' + ' '.join(attrs) if attrs else ''}>{icon}{html.escape(text)}</td>")
        out.append(f'<tr style="{rstyle}">' + "".join(cells_html) + "</tr>")
    out.append("</table></div></body></html>")
    return "".join(out), hits
