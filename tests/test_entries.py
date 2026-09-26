import io
import zipfile
from datetime import date, datetime

import openpyxl
import pytest

from mis_reports import Filters, load_master, run_reports
from mis_reports.entries import EntryError, MasterFile, restore_latest_backup
from mis_reports.loader import ORDER_COLUMNS
from mis_reports.xlsx_edit import Book, _Eval, relative_pattern, translate
from tests.test_reports import PUR_HEADERS, SAL_HEADERS, pur, sale

ORDER_HEADERS = ["Contract No", "Date", "Sales Contract No", "Company Name", "Party Name", "Branch", "Broker Name",
                 "Quality", "Qty Mts", "Rate", "Packing", "Delivery Date", "Delivery Place", "Cancel Qty",
                 "Revised Qty", "Recd Qty", "Qty + / -", "Bal Qty", "AMT", "Status", "Cancel date", "Recd Date",
                 "Bill Date", "Bill No", "BILL RATE", "LORRY NO"]


def _bill_sheet(wb, name, headers, rows, formulas):
    ws = wb.create_sheet(name)
    ws.append([None])
    ws.append([None])
    ws.append(headers)
    col = {h: openpyxl.utils.get_column_letter(i + 1) for i, h in enumerate(headers)}
    for i, r in enumerate(rows, 1):
        ws.append([i, i] + [r.get(h) for h in headers[2:]])
        n = ws.max_row
        for h, f in formulas.items():
            ws[f"{col[h]}{n}"] = "=" + f.format(**{k.replace(" ", "_").replace(".", ""): f"{v}{n}"
                                                  for k, v in col.items()})
            ws[f"{col[h]}{n}"].number_format = "0.00"
    for r in range(4, ws.max_row + 1):
        ws[f"{col['BILL DATE']}{r}"].number_format = "dd-mm-yyyy"
    return ws


def build_master(path):
    """A small workbook laid out like the Master Sheet, with the usual formulas."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _bill_sheet(wb, "Purchase", PUR_HEADERS,
                [pur("OLD SUPPLIER", "Urad", 1000, 90, datetime(2026, 8, 1)),
                 pur("OLD SUPPLIER", "Urad", 2000, 91, datetime(2026, 8, 5))],
                {"Amount befor GST": "{WEIGHT}*{RATE}", "Total Pre Bill Exp": "SUM({Adhat}:{Other})",
                 "Net Amount befor GST": "{Amount_befor_GST}+{Total_Pre_Bill_Exp}-{Discount}",
                 "GST": "{Net_Amount_befor_GST}*5%", "Bill Amt": "{Amount_befor_GST}+{GST}",
                 "Final Amt.": "{Bill_Amt}", "Total Post exp.": "{Storage_Amt}+SUM({Brokerage}:{Transportation})"})
    _bill_sheet(wb, "Sales", SAL_HEADERS, [sale("OLD BUYER", "Urad", 1000, 100, datetime(2026, 8, 10))],
                {"Amount before GST": "{WH_Weight}*{RATE}", "GST": "{Amount_before_GST}*5%",
                 "Bill Amt": "{Amount_before_GST}+{GST}", "Net Amount befor GST": "{Amount_before_GST}-{Discount}"})
    for name in ("PO", "SO"):
        ws = wb.create_sheet(name)
        ws.append([None])
        ws.append(ORDER_HEADERS)
        ws.append([1, datetime(2026, 7, 20), None, "SCPL", "OLD PARTY", None, None, "URAD", 50, 97.5, None, None,
                   None, 0, "=I3-N3", 0, 0, "=O3-P3+Q3", "=(R3*1000)*J3", '=IF(R3=0,"Closed","Open")'])
    ws = wb.create_sheet("Notes")
    ws["A1"] = "keep me"
    wb.save(path)
    return path


@pytest.fixture
def master(tmp_path):
    return MasterFile(build_master(tmp_path / "master.xlsx"))


def test_translate_and_patterns():
    assert translate("A4+1", 1) == "A5+1"
    assert translate("IF(AS4=\"Paid\",\"A1\",($AT$1+1)-M4)", 2) == "IF(AS6=\"Paid\",\"A1\",($AT$1+1)-M6)"
    assert translate("SUM(AC4:AI4)", 0, 1) == "SUM(AD4:AJ4)"
    assert translate("VLOOKUP(H4,masters!B:C,2,0)", 5) == "VLOOKUP(H9,masters!B:C,2,0)"
    assert relative_pattern("A7+B8*$C$1", 8) == "A[-1]+B[0]*$C$1"


def test_evaluator():
    cells = {"A1": 10.0, "B1": 4.0, "C1": "Paid", "D1": None}
    ev = lambda f: _Eval(f, cells.get).run()
    assert ev("A1*B1+5%") == pytest.approx(40.05)
    assert ev("ROUND(A1/3,2)") == 3.33
    assert ev('IF(C1="Paid","Closed",A1-B1)') == "Closed"
    assert ev('IF(D1="Paid","Closed",A1-B1)') == 6
    assert ev("SUM(A1:B1)-(B1*2)") == 6
    assert ev("ROUND(2.5,0)") == 3   # Excel rounds half away from zero


def test_shared_formula_is_unshared_before_overwrite(tmp_path):
    p = tmp_path / "s.xlsx"
    wb = openpyxl.Workbook()
    for i in range(1, 4):
        wb.active[f"A{i}"] = i
        wb.active[f"B{i}"] = f"=A{i}*2"
    wb.save(p)
    # turn B1:B3 into one shared formula, as Excel stores fill-down
    with zipfile.ZipFile(p) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    xml = parts["xl/worksheets/sheet1.xml"].decode()
    xml = xml.replace("<f>A1*2</f>", '<f t="shared" ref="B1:B3" si="0">A1*2</f>')
    xml = xml.replace("<f>A2*2</f>", '<f t="shared" si="0"/>').replace("<f>A3*2</f>", '<f t="shared" si="0"/>')
    parts["xl/worksheets/sheet1.xml"] = xml.encode()
    with zipfile.ZipFile(p, "w") as z:
        for n, d in parts.items():
            z.writestr(n, d)

    book = Book(p)
    sh = book.sheet("Sheet")
    assert sh.formula("B3") == "A3*2"
    sh.set("B1", 99)
    sh.recalc_row(3)
    assert sh.formula("B3") == "A3*2" and sh.value("B3") == 6
    book.save()
    ws = openpyxl.load_workbook(p)["Sheet"]
    assert ws["B1"].value == 99 and ws["B2"].value == "=A2*2" and ws["B3"].value == "=A3*2"


def test_full_flow_updates_sheet_and_reports(master):
    po = master.add_order("PO", {"date": date(2026, 9, 1), "party": "NEW SUPPLIER", "commodity": "URAD",
                                 "qty_mt": 10, "rate": 90})
    row = master.add_invoice("Purchase", {"party": "NEW SUPPLIER", "commodity": "Urad", "qty": 10_000, "rate": 90,
                                          "bill_date": date(2026, 9, 2), "stock_date": date(2026, 9, 2),
                                          "bill_no": "N-1", "type": "Local"}, gst_pct=5, order_row=po)
    master.add_note("Purchase", row, 1_000)
    master.set_expenses("Purchase", row, {"adhat": 2_000, "brokerage": 900})
    master.add_settlement("Purchase", date(2026, 9, 10), [(row, 500_000, 1_000)])
    so = master.add_order("SO", {"date": date(2026, 9, 3), "party": "NEW BUYER", "commodity": "URAD",
                                 "qty_mt": 10, "rate": 100})
    srow = master.add_invoice("Sales", {"party": "NEW BUYER", "commodity": "Urad", "qty": 10_000, "rate": 100,
                                        "bill_date": date(2026, 9, 4), "bill_no": "S-1"}, gst_pct=5, order_row=so)
    master.add_note("Sales", srow, 500)
    master.add_settlement("Sales", date(2026, 9, 12), [(srow, 1_049_500, 0)], outstanding={srow: 1_049_500})

    d = load_master(master.path)
    p = d.purchase.set_index("source_row").loc[row]
    assert p.amount == 900_000 and p.pre_exp == 2_000 and p.net_amount == 902_000
    assert p.gst == pytest.approx(45_100) and p.bill_amt == pytest.approx(945_100)
    assert p.paid == 500_000 and p.tds == 1_000 and p.note == 1_000 and p.post_exp == 900
    assert p.pay_date == pd_ts(2026, 9, 10)
    s = d.sales.set_index("source_row").loc[srow]
    assert s.bill_amt == pytest.approx(1_050_000) and s.status == "Recd"
    assert d.po.iloc[-1].bal_qty_mt == 0 and d.po.iloc[-1].status == "Closed"
    assert d.so.iloc[-1].contract_no == "2"

    rs = run_reports(None, Filters(as_of=date(2026, 9, 30)), data=d)
    cred = rs.creditors.summary.set_index("party").loc["NEW SUPPLIER"]
    assert cred.closing == pytest.approx(945_100 - 500_000 - 1_000 - 1_000)
    assert rs.debtors.summary.set_index("party").loc["NEW BUYER"].closing == pytest.approx(0)

    # untouched sheet survives, a backup exists for every save and the log lists every entry
    assert openpyxl.load_workbook(master.path)["Notes"]["A1"].value == "keep me"
    assert len(list(master.backup_dir.iterdir())) == 9
    assert [r["Action"] for r in master.read_log()][::-1] == [
        "Purchase Order", "Purchase Invoice", "Debit Note", "Expenses", "Payment", "Sales Order",
        "Sales Invoice", "Credit Note", "Receipt"]


def test_undo_restores_previous_version(master):
    before = master.path.read_bytes()
    master.add_order("PO", {"date": date(2026, 9, 1), "party": "X", "commodity": "URAD", "qty_mt": 1, "rate": 1})
    assert master.path.read_bytes() != before
    assert restore_latest_backup(master.path, master.backup_dir)
    assert master.path.read_bytes() == before


def test_validation_errors(master):
    with pytest.raises(EntryError, match="Party"):
        master.add_invoice("Purchase", {"commodity": "Urad", "qty": 1, "rate": 1, "bill_date": date.today()})
    with pytest.raises(EntryError):
        master.add_note("Purchase", 999, 10)
    with pytest.raises(EntryError):
        master.add_settlement("Purchase", date.today(), [(4, 0, 0)])


def pd_ts(y, m, d):
    import pandas as pd
    return pd.Timestamp(y, m, d)


# --------------------------------------------------------------------------- numbering, multi-item bills, notes
def test_next_number_follows_the_series():
    from mis_reports.entries import next_number
    assert next_number([1, 2, 3]) == 4
    assert next_number([]) == 1
    assert next_number(["SCPL/CH12/26-27", "SCPL/CH13/26-27"]) == "SCPL/CH14/26-27"
    assert next_number(["SCPL/CH13/26-27"]) == "SCPL/CH14/26-27"   # not the year
    assert next_number(["PO/26-27/009"]) == "PO/26-27/010"
    assert next_number([], "DN-0001") == "DN-0001"
    assert next_number(["DN-0001", "DN-0002"], "DN-0001") == "DN-0003"


def test_order_numbering_auto_and_manual(master):
    master.add_order("PO", {"date": date(2026, 9, 1), "party": "A", "commodity": "URAD", "qty_mt": 1, "rate": 1})
    master.add_order("PO", {"date": date(2026, 9, 1), "party": "A", "commodity": "URAD", "qty_mt": 1, "rate": 1},
                     contract_no="PO/26-27/007")
    with pytest.raises(EntryError, match="already used"):
        master.add_order("PO", {"date": date(2026, 9, 1), "party": "A", "commodity": "URAD", "qty_mt": 1,
                                "rate": 1}, contract_no="po/26-27/007")
    d = load_master(master.path)
    assert d.po["contract_no"].tolist() == ["1", "2", "PO/26-27/007"]
    assert master.next_contract_no("PO") == "PO/26-27/008"


def test_multi_item_invoice_with_contracts_and_pre_bill_expenses(master):
    po1 = master.add_order("PO", {"date": date(2026, 9, 1), "party": "S", "commodity": "URAD", "qty_mt": 2, "rate": 90})
    po2 = master.add_order("PO", {"date": date(2026, 9, 1), "party": "S", "commodity": "TOOR", "qty_mt": 1, "rate": 100})
    rows = master.add_invoice_lines(
        "Purchase", {"party": "S", "bill_no": "B-9", "bill_date": date(2026, 9, 5)},
        [{"commodity": "Urad", "qty": 2_000, "rate": 90, "gst_pct": 5, "order_row": po1},
         {"commodity": "Toor", "qty": 1_000, "rate": 100, "gst_pct": 0, "order_row": po2}],
        expenses={"adhat": 1_400, "discount": 280})
    d = load_master(master.path)
    p = d.purchase.set_index("source_row").loc[rows]
    assert p["adhat"].tolist() == [900, 500]            # split 180,000 : 100,000
    assert p["discount"].tolist() == [180, 100]
    assert p["gst"].tolist() == [pytest.approx(180_720 * 0.05), 0]
    assert (d.po.set_index("source_row").loc[[po1, po2], "status"] == "Closed").all()
    assert d.po.set_index("source_row").loc[po1, "bill_no"] == "B-9"


def test_detailed_notes_flow_into_pnl_stock_and_ledgers(master):
    row = master.add_invoice("Purchase", {"party": "S", "commodity": "Urad", "qty": 1_000, "rate": 100,
                                          "bill_date": date(2026, 9, 1), "bill_no": "P-1"}, gst_pct=5)
    srow = master.add_invoice_lines("Sales", {"party": "C", "bill_no": "S-1", "bill_date": date(2026, 9, 2)},
                                    [{"commodity": "Urad", "qty": 500, "rate": 120, "gst_pct": 5, "link_row": row}])[0]
    assert master.add_note_detailed("Purchase", row, {"qty": 100, "rate": 100, "gst_pct": 5, "other": 50,
                                                      "reason": "Goods returned"}) == "DN-0001"
    master.add_note_detailed("Sales", srow, {"qty": 50, "rate": 120, "gst_pct": 5, "note_no": "MY-CN-1"})
    with pytest.raises(EntryError, match="more than"):
        master.add_note_detailed("Purchase", row, {"qty": 5_000, "rate": 1})

    d = load_master(master.path)
    assert d.debit_notes.iloc[0][["qty", "taxable", "gst", "other", "total"]].tolist() == [100, 10_000, 500, 50, 10_550]
    assert d.credit_notes.iloc[0]["note_no"] == "MY-CN-1"
    assert d.purchase.set_index("source_row").loc[row, "sales_inv"] == "S-1"
    assert d.sales.set_index("source_row").loc[srow, "purchase_inv"] == "P-1"

    rs = run_reports(None, Filters(as_of=date(2026, 9, 30)), data=d)
    urad = rs.pnl.set_index("product").loc["URAD"]
    assert urad.net_purchase_qty == 900 + 3_000       # existing sample rows hold 3,000 kg
    assert urad.purchase_returns == 10_050
    assert urad.sales_return_qty == 50 and urad.sales_returns == 6_000
    stock = rs.stock.set_index("product").loc["URAD"]
    assert stock.closing_qty == (3_000 + 1_000 - 100) - (1_000 + 500 - 50)
    cred = rs.creditors.summary.set_index("party").loc["S"]
    assert cred.note == pytest.approx(10_550)
    m = rs.margins.iloc[0]
    assert m.purchase_bill == "P-1" and m.cost_rate == 100 and m.margin == pytest.approx(500 * 20)
    assert len(rs.po_detail) == 1 and rs.po_detail.iloc[0].days_open == 72
    wb = openpyxl.load_workbook(master.path)
    assert wb.sheetnames[-2:] == ["Debit Notes", "Credit Notes"] and wb["Notes"]["A1"].value == "keep me"
