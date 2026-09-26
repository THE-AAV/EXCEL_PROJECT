from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill

from ui.sheet_viewer import format_value, load_workbook_views, render_html

ACC = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'


def test_number_and_date_formats():
    assert format_value(2230340, ACC) == "22,30,340.00"
    assert format_value(2230340, ACC, indian=False) == "2,230,340.00"
    assert format_value(-1234.5, ACC) == "(1,234.50)"
    assert format_value(0, ACC) == "-"
    assert format_value(0.1234, "0.0%") == "12.3%"
    assert format_value(25060, "General") == "25060"
    assert format_value(12.5, '"₹"#,##0.00') == "₹12.50"
    assert format_value(datetime(2026, 8, 13), "mm-dd-yy") == "08-13-26"
    assert format_value(datetime(2026, 8, 13), "dd-mmm-yyyy") == "13-Aug-2026"
    assert format_value(True, "General") == "TRUE"


def test_workbook_view(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"], ws["B1"] = "Item", "Amount"
    ws["A1"].font = Font(bold=True)
    ws["B1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["A2"], ws["B2"] = "Urad", 1000
    ws["B3"] = "=B2*2"
    ws.merge_cells("A4:B4")
    ws["A4"] = "Merged"
    ws.freeze_panes = "A2"
    wb.create_sheet("Second")["C3"] = "x"
    wb.active = 0
    p = tmp_path / "v.xlsx"
    wb.save(p)

    views = load_workbook_views(str(p))
    assert [v["name"] for v in views] == ["Data", "Second"] and views[0]["active"]
    v = views[0]
    assert v["freeze"] == (1, 0) and v["merged"][(4, 1)] == (1, 2)
    assert v["cells"][(3, 2)]["f"] == "=B2*2" and v["cells"][(3, 2)]["p"]   # not calculated yet
    html, hits = render_html(v, find="urad")
    assert hits == 1 and "Urad" in html and "Merged" not in html   # only matching rows (+ frozen header)
    html, _ = render_html(v)
    assert "Merged" in html and 'colspan="2"' in html
    assert "background:#FFFF00" in html and "font-weight:700" in html
    html, _ = render_html(v, show_formulas=True)
    assert "=B2*2" in html


def test_filter_and_sort(tmp_path):
    from ui.sheet_viewer import BLANK, column_values, filter_rows, header_labels, is_numeric_column
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Party", "Qty", "Date"])
    for row in (["A", 10, datetime(2026, 8, 1)], ["B", 5, datetime(2026, 8, 3)], ["A", 30, datetime(2026, 8, 2)],
                [None, 7, None]):
        ws.append(row)
    ws.auto_filter.ref = "A1:C5"
    p = tmp_path / "f.xlsx"
    wb.save(p)
    v = load_workbook_views(str(p))[0]
    assert v["autofilter"] == (1, 1, 3)
    assert header_labels(v, 1) == {1: "A · Party", 2: "B · Qty", 3: "C · Date"}
    assert column_values(v, 1, 1) == ["A", "B", BLANK]
    assert column_values(v, 1, 2) == ["5", "7", "10", "30"]          # numbers in number order
    assert is_numeric_column(v, 1, 2) and not is_numeric_column(v, 1, 1)
    assert filter_rows(v, 1, {1: {"values": ["A"]}}) == [2, 4]
    assert filter_rows(v, 1, {1: {"values": [BLANK]}}) == [5]
    assert filter_rows(v, 1, {2: {"min": 7, "max": 20}}) == [2, 5]
    assert filter_rows(v, 1, {}, sort=(2, False)) == [4, 2, 5, 3]
    assert filter_rows(v, 1, {}, sort=(3, True)) == [2, 4, 3, 5]      # blanks last
    html, _ = render_html(v, header_row=1, shown_rows=[4, 2], filter_cols={1})
    assert html.index(">30<") < html.index(">10<") and 'class="fb on"' in html
