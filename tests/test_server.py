"""The web backend: sign-in, roles, Master Sheet import and reports from the database.

Runs against SQLite by default; set TEST_DATABASE_URL (e.g. postgresql+psycopg://user@localhost/test_db)
to run the same tests on PostgreSQL.
"""
import io
import os
from datetime import date

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, text

from mis_reports import Filters, load_master
from mis_reports.entries import MasterFile
from mis_reports.pipeline import build_reports
from server import db
from server.config import Settings
from server.main import create_app
from server.reports_api import report_json
from tests.test_entries import build_master

AS_OF = date(2026, 9, 30)
H = {"X-Requested-With": "fetch"}
PASSWORD = "correct-horse-1"


@pytest.fixture
def settings(tmp_path):
    url = os.environ.get("TEST_DATABASE_URL")
    if url:  # start every test from an empty database
        eng = create_engine(url)
        db.metadata.drop_all(eng)
        eng.dispose()
    return Settings(data_dir=tmp_path / "data", database_url=url or f"sqlite:///{tmp_path / 'test.db'}",
                    secret_key="test-secret-" + "x" * 32)


@pytest.fixture
def client(settings):
    with TestClient(create_app(settings)) as c:
        c.headers.update(H)
        yield c


@pytest.fixture
def rich_master(tmp_path):
    """A Master Sheet with orders, linked invoices and debit / credit notes, made with the entry code."""
    m = MasterFile(build_master(tmp_path / "master.xlsx"))
    row = m.add_invoice("Purchase", {"party": "S", "commodity": "Urad", "qty": 1_000, "rate": 100,
                                     "bill_date": date(2026, 9, 1), "bill_no": "P-1"}, gst_pct=5)
    srow = m.add_invoice_lines("Sales", {"party": "C", "bill_no": "S-1", "bill_date": date(2026, 9, 2)},
                               [{"commodity": "Urad", "qty": 500, "rate": 120, "gst_pct": 5, "link_row": row}])[0]
    m.add_note_detailed("Purchase", row, {"qty": 100, "rate": 100, "gst_pct": 5, "other": 50,
                                          "reason": "Goods returned"})
    m.add_note_detailed("Sales", srow, {"qty": 50, "rate": 120, "gst_pct": 5})
    return m.path


def assert_same(a, b, where=""):
    """Deep equality, with numbers equal to 9 significant digits (JSON turns whole floats into ints)."""
    if isinstance(a, dict) and isinstance(b, dict):
        assert a.keys() == b.keys(), where
        for k in a:
            assert_same(a[k], b[k], f"{where}.{k}")
    elif isinstance(a, list) and isinstance(b, list):
        assert len(a) == len(b), where
        for i, (x, y) in enumerate(zip(a, b)):
            assert_same(x, y, f"{where}[{i}]")
    elif isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        assert a == pytest.approx(b, rel=1e-9, abs=1e-9), where
    else:
        assert a == b, where


def setup_admin(client):
    r = client.post("/api/setup", json={"username": "Owner", "full_name": "Owner", "role": "admin",
                                        "password": PASSWORD})
    assert r.status_code == 200, r.text
    return r


def add_user(client, name, role):
    r = client.post("/api/users", json={"username": name, "role": role, "password": PASSWORD})
    assert r.status_code == 201, r.text


def sign_in(client, name):
    client.cookies.clear()
    r = client.post("/api/auth/login", json={"username": name, "password": PASSWORD})
    assert r.status_code == 200, r.text
    return r.json()


def upload(client, path):
    with open(path, "rb") as fh:
        return client.post("/api/imports", files={"file": ("Master Sheet.xlsx", fh.read())})


def test_first_admin_setup_only_once(client):
    assert client.get("/api/setup").json() == {"needs_first_admin": True}
    assert setup_admin(client).json()["role"] == "admin"
    assert client.get("/api/auth/me").json()["username"] == "owner"
    assert client.get("/api/setup").json() == {"needs_first_admin": False}
    r = client.post("/api/setup", json={"username": "intruder", "role": "admin", "password": PASSWORD})
    assert r.status_code == 409


def test_sign_in_is_required_and_wrong_passwords_lock_the_account(client):
    setup_admin(client)
    client.cookies.clear()
    assert client.get("/api/reports").status_code == 401
    for _ in range(5):
        r = client.post("/api/auth/login", json={"username": "owner", "password": "wrong-password"})
        assert r.status_code == 401
    r = client.post("/api/auth/login", json={"username": "owner", "password": PASSWORD})
    assert r.status_code == 423 and "Try again" in r.json()["detail"]


def test_changes_need_the_page_header(client):
    setup_admin(client)
    r = client.post("/api/auth/logout", headers={"X-Requested-With": ""})
    assert r.status_code == 403


def test_roles_are_enforced_by_the_backend(client, rich_master):
    setup_admin(client)
    add_user(client, "clerk", "editor")
    add_user(client, "partner", "viewer")

    sign_in(client, "partner")
    assert client.get("/api/reports").status_code == 200
    assert client.get("/api/aliases").status_code == 200
    assert upload(client, rich_master).status_code == 403
    assert client.get("/api/users").status_code == 403
    assert client.put("/api/aliases", json=[]).status_code == 403
    assert client.get("/api/audit").status_code == 403

    sign_in(client, "clerk")
    assert client.get("/api/reports").status_code == 200
    assert upload(client, rich_master).status_code == 403
    assert client.post("/api/users", json={"username": "x", "role": "admin", "password": PASSWORD}).status_code == 403


def test_switching_a_user_off_ends_their_session(client):
    setup_admin(client)
    add_user(client, "clerk", "editor")
    uid = next(u["id"] for u in client.get("/api/users").json() if u["username"] == "clerk")
    admin_cookie = client.cookies.get("session")
    sign_in(client, "clerk")
    clerk_cookie = client.cookies.get("session")

    client.cookies.set("session", admin_cookie)
    assert client.patch(f"/api/users/{uid}", json={"active": False}).json()["active"] is False
    client.cookies.set("session", clerk_cookie)
    assert client.get("/api/auth/me").status_code == 401
    r = client.post("/api/auth/login", json={"username": "clerk", "password": PASSWORD})
    assert r.status_code == 403


def test_last_admin_cannot_be_removed(client):
    setup_admin(client)
    me = client.get("/api/auth/me").json()
    assert client.patch(f"/api/users/{me['id']}", json={"role": "viewer"}).status_code == 409
    assert client.patch(f"/api/users/{me['id']}", json={"active": False}).status_code == 409


def test_import_preview_saves_nothing_until_confirmed(client, rich_master):
    setup_admin(client)
    r = upload(client, rich_master)
    assert r.status_code == 201, r.text
    batch = r.json()
    assert batch["status"] == "pending"
    s = batch["summary"]
    assert s["counts"]["purchase"] == 3 and s["counts"]["sales"] == 2 and s["counts"]["debit_notes"] == 2
    assert s["counts"]["po"] == 1 and s["counts"]["credit_notes"] == 1
    assert s["new_parties"] == ["C", "OLD BUYER", "OLD SUPPLIER", "S"] and s["new_products"] == ["URAD"]
    assert s["current_counts"] is None

    assert client.get("/api/status").json()["has_data"] is False
    assert client.get("/api/reports", params={"as_of": AS_OF}).json()["pnl"] == []

    r = client.post(f"/api/imports/{batch['id']}/confirm")
    assert r.status_code == 200 and r.json()["status"] == "active"
    assert client.post(f"/api/imports/{batch['id']}/confirm").status_code == 409
    st = client.get("/api/status").json()
    assert st["has_data"] and st["source"]["filename"] == "Master Sheet.xlsx"
    assert st["options"]["products"] == ["URAD"]


def test_reports_from_the_database_match_the_excel_file(client, rich_master):
    """Phase 1's check: every report worked out from the database equals the one from the Master Sheet file."""
    setup_admin(client)
    client.post(f"/api/imports/{upload(client, rich_master).json()['id']}/confirm")

    for f in (Filters(as_of=AS_OF), Filters(as_of=date(2026, 8, 31)),
              Filters(as_of=AS_OF, date_from=date(2026, 9, 1), date_to=date(2026, 9, 30)),
              Filters(as_of=AS_OF, products=["URAD"], companies=["SCPL"])):
        params = {"as_of": f.as_of, "date_from": f.date_from, "date_to": f.date_to,
                  "products": f.products, "companies": f.companies}
        from_db = client.get("/api/reports", params={k: v for k, v in params.items() if v}).json()
        from_file = report_json(build_reports(load_master(rich_master), Filters(**{**f.__dict__, "aliases": {}})))
        assert_same(from_db, from_file, str(f))

    body = client.get("/api/reports", params={"as_of": AS_OF}).json()
    assert body["dashboard"]["purchase_bills"] == 3
    assert body["creditors"]["summary"] and body["debit_notes"] and body["margins"]


def test_excel_download_and_product_names(client, rich_master):
    setup_admin(client)
    client.post(f"/api/imports/{upload(client, rich_master).json()['id']}/confirm")
    r = client.get("/api/reports/excel", params={"as_of": AS_OF})
    assert r.status_code == 200
    assert r.headers["content-disposition"] == 'attachment; filename="Business_Reports_2026-09-30.xlsx"'
    assert {"Summary", "Product P&L"} <= set(openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames)

    saved = client.put("/api/aliases", json=[{"name_in_sheet": "Urad", "report_as": "Black Gram"}]).json()
    assert saved == [{"name_in_sheet": "Urad", "report_as": "Black Gram"}]
    pnl = client.get("/api/reports", params={"as_of": AS_OF}).json()["pnl"]
    assert [p["product"] for p in pnl] == ["BLACK GRAM"]


def test_a_later_import_replaces_the_data_and_the_earlier_one_can_be_restored(client, rich_master, tmp_path):
    setup_admin(client)
    first = upload(client, rich_master).json()["id"]
    client.post(f"/api/imports/{first}/confirm")

    small = build_master(tmp_path / "small.xlsx")
    second = upload(client, small).json()
    assert second["summary"]["current_counts"]["purchase"] == 3
    assert second["summary"]["new_parties"] == []
    client.post(f"/api/imports/{second['id']}/confirm")
    assert client.get("/api/reports", params={"as_of": AS_OF}).json()["dashboard"]["purchase_bills"] == 2

    statuses = {b["id"]: b["status"] for b in client.get("/api/imports").json()}
    assert statuses == {first: "replaced", second["id"]: "active"}
    assert client.post(f"/api/imports/{first}/restore").status_code == 200
    assert client.get("/api/reports", params={"as_of": AS_OF}).json()["dashboard"]["purchase_bills"] == 3

    actions = [a["action"] for a in client.get("/api/audit").json()]
    assert actions[:3] == ["import.restore", "import.confirm", "import.preview"]


def test_bad_files_are_refused_and_discard_removes_a_preview(client, rich_master, tmp_path):
    setup_admin(client)
    r = client.post("/api/imports", files={"file": ("notes.txt", b"hello")})
    assert r.status_code == 422
    r = client.post("/api/imports", files={"file": ("fake.xlsx", b"not a workbook")})
    assert r.status_code == 422 and "could not be opened" in r.json()["detail"]
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "empty.xlsx")
    r = upload(client, tmp_path / "empty.xlsx")
    assert r.status_code == 422 and "Purchase" in r.json()["detail"]

    batch = upload(client, rich_master).json()
    assert client.post(f"/api/imports/{batch['id']}/discard").json() == {"ok": True}
    assert client.post(f"/api/imports/{batch['id']}/confirm").status_code == 409


def test_workbook_without_notes_imports(client, tmp_path):
    """Lesson from the Streamlit app: workbooks without debit / credit note sheets must still work."""
    setup_admin(client)
    client.post(f"/api/imports/{upload(client, build_master(tmp_path / 'plain.xlsx')).json()['id']}/confirm")
    body = client.get("/api/reports", params={"as_of": AS_OF}).json()
    assert body["pnl"][0]["product"] == "URAD" and body["debit_notes"] == []


def test_password_change_signs_out_other_sessions(client):
    setup_admin(client)
    old = client.cookies.get("session")
    r = client.post("/api/auth/password", json={"current_password": PASSWORD, "new_password": "another-pass-2"})
    assert r.status_code == 200
    assert client.get("/api/auth/me").status_code == 200
    client.cookies.set("session", old)
    assert client.get("/api/auth/me").status_code == 401


def test_health(client, settings):
    assert client.get("/api/health").json() == {"ok": True}
    eng = db.make_engine(settings.database_url)
    with eng.connect() as conn:
        assert conn.execute(text("select count(*) from users")).scalar_one() == 0
