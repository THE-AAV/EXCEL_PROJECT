import io
from datetime import date, datetime

import openpyxl
import pytest

from mis_reports import Filters, run_reports
from mis_reports.loader import PURCHASE_COLUMNS, SALES_COLUMNS, load_master

PUR_HEADERS = ["SR.NO.", "PSN", *PURCHASE_COLUMNS.values()]
SAL_HEADERS = ["SR.NO.", "PSN", *SALES_COLUMNS.values()]


def _sheet(wb, name, headers, rows):
    ws = wb.create_sheet(name)
    ws.append([None])
    ws.append([None])
    ws.append(headers)  # header on row 3, like the Master Sheet
    for i, r in enumerate(rows, 1):
        ws.append([i, i] + [r.get(h) for h in headers[2:]])


def make_master(purchase, sales, opening=None) -> io.BytesIO:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Purchase", PUR_HEADERS, purchase)
    _sheet(wb, "Sales", SAL_HEADERS, sales)
    ws = wb.create_sheet("Total Debtor Creditor")
    ws.append([None])
    ws.append([None])
    ws.append(["PSN", "Party Name", "Opening Bal."])
    for name, bal in (opening or {}).items():
        ws.append([1, name, bal])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def pur(party, commodity, qty, rate, when, **kw):
    amt = qty * rate
    row = {"PARTY NAME": party, "COMMODITY": commodity, "WEIGHT": qty, "RATE": rate,
           "Amount befor GST": amt, "BILL DATE": when, "Type": "Local", "Com.": "SCPL",
           "GST": amt * 0.05, "BILL NO.": f"P-{party}-{when:%d%m}"}
    row.update(kw)
    return row


def sale(party, commodity, qty, rate, when, **kw):
    amt = qty * rate
    row = {"PARTY NAME": party, "COMMODITY": commodity, "WH Weight": qty, "RATE": rate,
           "Amount before GST": amt, "BILL DATE": when, "Type": "Local", "Com": "SCPL",
           "GST": amt * 0.05, "BILL NO.": f"S-{party}-{when:%d%m}"}
    row.update(kw)
    return row


D1, D2, D3 = datetime(2026, 5, 1), datetime(2026, 7, 1), datetime(2026, 9, 1)
AS_OF = date(2026, 9, 26)


def test_pnl_matches_display_sheet_when_everything_is_sold():
    src = make_master(
        [pur("A", "Urad", 1000, 90, D1, Adhat=500, Discount=200, **{"Total Post exp.": 300})],
        [sale("X", "urad ", 1000, 100, D2, Discount=1000, **{"Total Post Exp.": 100})])
    p = run_reports(src, Filters(as_of=AS_OF)).pnl.set_index("product").loc["URAD"]
    assert p.net_purchase == pytest.approx(90_000 + 500 - 200)
    assert p.net_sales == pytest.approx(99_000)
    assert p.closing_qty == 0
    # Display sheet: Total Sales - (Total Purchase + post-bill expenses)
    assert p.net_profit == pytest.approx(99_000 - 90_300 - 300 - 100)


def test_unsold_stock_is_carried_at_average_cost():
    src = make_master([pur("A", "Toor", 1000, 50, D1), pur("B", "Toor", 1000, 70, D2)],
                      [sale("X", "Toor", 500, 80, D3)])
    p = run_reports(src, Filters(as_of=AS_OF)).pnl.set_index("product").loc["TOOR"]
    assert p.avg_cost_rate == pytest.approx(60)
    assert p.closing_value == pytest.approx(1500 * 60)
    assert p.cogs == pytest.approx(500 * 60)
    assert p.net_profit == pytest.approx(500 * 20)


def test_product_aliases_and_period_filter():
    src = make_master([pur("A", "SOYA BEANS", 100, 40, D1), pur("A", "Soyabean", 100, 40, D3)],
                      [sale("X", "Soyabean", 100, 50, D1)])
    rs = run_reports(src, Filters(as_of=AS_OF, aliases={"soya beans": "SOYABEAN"}))
    assert rs.pnl["product"].tolist() == ["SOYABEAN"]
    assert rs.pnl.iloc[0].purchase_qty == 200
    src.seek(0)
    rs = run_reports(src, Filters(as_of=AS_OF, date_to=date(2026, 6, 30), aliases={"SOYA BEANS": "SOYABEAN"}))
    assert rs.pnl.iloc[0].purchase_qty == 100


def test_creditor_fifo_ageing_opening_and_party_spellings():
    src = make_master(
        [pur("R.R. Impex", "Urad", 100, 100, D1, **{"Final Amt.": 10_000, "Amount": 12_000, "TDS": 10}),
         pur("R R IMPEX", "Urad", 100, 100, D3, **{"Final Amt.": 10_000})],
        [], opening={"RR Impex": 5_000})
    led = run_reports(src, Filters(as_of=AS_OF)).creditors
    assert len(led.summary) == 1
    s = led.summary.iloc[0]
    assert s.opening == 5_000 and s.billed == 20_000 and s.paid == 12_000
    assert s.closing == pytest.approx(5_000 + 20_000 - 12_000 - 10)
    # 12,010 settled: clears the 5,000 opening, then 7,010 off the May bill
    assert s["Over 90 days"] == pytest.approx(10_000 - 7_010)
    assert s["0-30 days"] == pytest.approx(10_000)
    assert led.bills["outstanding"].sum() == pytest.approx(s.closing)


def test_debtors_ignore_payments_after_as_of_date():
    src = make_master([], [sale("X", "Urad", 100, 100, D1, **{"Bill Amt": 10_500, "Amount": 10_500,
                                                               "Date": datetime(2026, 10, 5)})])
    s = run_reports(src, Filters(as_of=AS_OF)).debtors.summary.iloc[0]
    assert s.paid == 0 and s.closing == pytest.approx(10_500)


def test_excel_export_has_all_sheets():
    src = make_master([pur("A", "Urad", 1000, 90, D1)], [sale("X", "Urad", 400, 100, D2)])
    wb = openpyxl.load_workbook(io.BytesIO(run_reports(src, Filters(as_of=AS_OF)).to_excel()))
    assert wb.sheetnames == ["Summary", "Product P&L", "Stock Summary", "Creditors", "Creditor Bills", "Debtors",
                             "Debtor Bills", "Purchase Orders", "Sales Orders", "Open Orders", "Debit Notes",
                             "Credit Notes", "Bill-wise Margin", "Data Checks"]
    assert wb["Product P&L"]["B4"].value == "URAD"


def test_missing_sheet_gives_clear_error():
    wb = openpyxl.Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    with pytest.raises(ValueError, match="Purchase"):
        load_master(buf)


def test_party_with_only_opening_balance_is_listed():
    src = make_master([pur("A", "Urad", 10, 10, D1, **{"Final Amt.": 100})], [], opening={"Old Supplier": 2_500})
    s = run_reports(src, Filters(as_of=AS_OF)).creditors.summary.set_index("party")
    assert s.loc["Old Supplier", "closing"] == 2_500
    assert s.loc["Old Supplier", "Undated"] == 2_500
